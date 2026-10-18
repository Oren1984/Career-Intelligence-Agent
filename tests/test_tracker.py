"""Tracker: vocabulary, Excel import/export round-trip, CRUD validation, workbook KPI definitions.

All workbook fixtures are synthetic and built here; the personal workbook is never read by tests.
"""
import datetime as dt
import io

import pytest
from openpyxl import Workbook, load_workbook

from app.db.models import Application
from app.tracker import service
from app.tracker.excel import COLUMNS, export_workbook, parse_date, preview_workbook
from app.tracker.vocab import SOURCES, STATUSES, TRACKS, label, to_code

HE_HEADERS = [pair[0] for pair in COLUMNS.values()]


def _workbook(rows, extra_header=None) -> bytes:
    """Synthetic workbook laid out like Job_Search_Tracker.xlsx (title rows, header row 4)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "מעקב משרות"
    ws["A1"] = "מעקב הגשת משרות - דוגמה"
    ws["A2"] = "שורת כותרת משנה"
    headers = HE_HEADERS + ([extra_header] if extra_header else [])
    for c, h in enumerate(headers, start=1):
        ws.cell(row=4, column=c, value=h)
    for r, row in enumerate(rows, start=5):
        for c, v in enumerate(row, start=1):
            ws.cell(row=r, column=c, value=v)
    lists = wb.create_sheet("רשימות בחירה")
    lists["A1"] = "סטטוסים (Status)"
    for i, code in enumerate(STATUSES, start=2):
        lists.cell(row=i, column=1, value=STATUSES[code][0])
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _row(ref, date, company, title, track="DevOps & ענן", source="לינקדאין (LinkedIn)", location="תל אביב · היברידי",
         salary="25-30K", status="הוגש", notes="", closure=""):
    days_formula = '=IF(B5="","",TODAY()-B5)'
    return [ref, date, company, title, track, source, location, salary, status, days_formula, notes, closure]


@pytest.fixture
def ws_session(db_session):
    db_session.info["workspace"] = "personal"
    return db_session


# ── Vocabulary ───────────────────────────────────────────────────────────────

def test_vocab_matches_workbook_lists():
    assert [v[0] for v in STATUSES.values()] == [
        "הוגש", "בבדיקה", "שיחה ראשונית", "ראיון מקצועי", "מטלת בית", "הצעה / מתקדם", "נדחה", "נסגר ללא מענה"]
    assert len(TRACKS) == 5 and len(SOURCES) == 14
    assert SOURCES["meta_ai"][0] == "Meta Al"   # original spelling preserved for round-trips


@pytest.mark.parametrize("field,value,code", [
    ("status", "הוגש", "submitted"), ("status", " Professional  interview ", "interview"),
    ("status", "offer", "offer"), ("source", "ג׳ובמאסטר (JobMaster)", "jobmaster"),
    ("source", "Meta AI", "meta_ai"), ("source", "meta al", "meta_ai"), ("track", "DevOps & Cloud", "devops_cloud"),
    ("status", "something else", None), ("status", "", None), ("status", None, None),
])
def test_to_code(field, value, code):
    assert to_code(field, value) == code


@pytest.mark.parametrize("value,expected,ok", [
    (dt.datetime(2026, 3, 4, 10, 0), dt.date(2026, 3, 4), True),
    ("04/03/2026", dt.date(2026, 3, 4), True),
    ("2026-03-04", dt.date(2026, 3, 4), True),
    (46085, dt.date(2026, 3, 4), True),           # Excel serial number
    ("", None, True), (None, None, True),
    ("next tuesday", None, False), ("31/02/2026", None, False),
])
def test_parse_date(value, expected, ok):
    assert parse_date(value) == (expected, ok)


# ── Import preview ───────────────────────────────────────────────────────────

def _mixed_workbook():
    return _workbook([
        _row(1, dt.datetime(2026, 9, 1), "Alpha Ltd", "DevOps Engineer", status="ראיון מקצועי"),
        _row(2, "15/09/2026", "Beta Inc", "Platform Engineer", source="Meta Al", status="מטלת בית"),
        _row(3, 46085, "Gamma", "SRE", status="סטטוס חדש לא מוכר", notes="note"),
        _row(4, "sometime", "Delta", "QA Lead", track="QA & אוטומציה"),
        _row(5, dt.datetime(2026, 9, 1), "alpha ltd ", "devops engineer"),          # duplicate of row 1
        [6, None, None, None, None, None, None, None, None, '=IF(B10="","",TODAY()-B10)', None, None],  # blank template
        [None, None, None, None, None, None, None, None, None, None, None, "orphan closure note"],      # incomplete
    ], extra_header="איש קשר")


def test_preview_classifies_rows_without_losing_data():
    p = preview_workbook(_mixed_workbook())
    assert p.error == "" and p.sheet == "מעקב משרות"
    assert len(p.rows) == 4
    assert p.blank_rows == 1
    assert len(p.incomplete_rows) == 1 and p.count("incomplete") == 1
    assert p.count("duplicate") == 1
    assert p.count("invalid_date") == 1
    assert p.count("unmapped") == 1
    assert p.extra_columns == ["איש קשר"]
    by_company = {r["company"]: r for r in p.rows}
    assert by_company["Alpha Ltd"]["status"] == "interview"
    assert by_company["Beta Inc"]["source"] == "meta_ai"
    assert by_company["Beta Inc"]["submitted_date"] == dt.date(2026, 9, 15)
    assert by_company["Gamma"]["status"] is None
    assert by_company["Gamma"]["extra"]["unmapped"]["status"] == "סטטוס חדש לא מוכר"
    assert by_company["Delta"]["submitted_date"] is None
    assert by_company["Delta"]["extra"]["unmapped"]["submitted_date"] == "sometime"
    # issues expose row numbers / fields only, never cell content
    assert all(set(vars(i)) == {"row", "kind", "field"} for i in p.issues)


def test_preview_skips_rows_already_saved(ws_session):
    p = preview_workbook(_mixed_workbook())
    service.commit_import(ws_session, p)
    again = preview_workbook(_mixed_workbook(), existing_keys=service.existing_keys(ws_session))
    assert again.rows == [] and again.count("duplicate") == 5


def test_preview_rejects_non_workbooks_and_missing_headers():
    assert preview_workbook(b"not an excel file").error
    wb = Workbook()
    wb.active["A1"] = "nothing here"
    buf = io.BytesIO()
    wb.save(buf)
    assert "header" in preview_workbook(buf.getvalue()).error


# ── Round-trip ───────────────────────────────────────────────────────────────

def _comparable(apps):
    return sorted(
        (a.submitted_date or dt.date.min, a.company, a.title, a.track, a.source, a.location, a.salary_range, a.status,
         a.notes, a.closure_reason, repr(sorted((a.extra or {}).items())))
        for a in apps
    )


@pytest.mark.parametrize("lang", ["he", "en"])
def test_excel_round_trip_preserves_everything(db_session, lang):
    db_session.info["workspace"] = "personal"
    service.commit_import(db_session, preview_workbook(_mixed_workbook()), include_incomplete=True)
    original = service.list_applications(db_session)

    exported = export_workbook(original, lang=lang)
    wb = load_workbook(io.BytesIO(exported))
    assert len(wb.sheetnames) == 3
    tracker = wb.worksheets[0]
    assert tracker.cell(row=4, column=3).value == COLUMNS["company"][0 if lang == "he" else 1]
    assert tracker.cell(row=5, column=10).value.startswith("=IF(B5")
    assert len(tracker.data_validations.dataValidation) == 3
    dash = wb.worksheets[2]
    assert dash["A5"].value.startswith("=COUNTIF(") and "TODAY" not in dash["A5"].value

    db_session.query(Application).delete()
    db_session.commit()
    service.commit_import(db_session, preview_workbook(exported), include_incomplete=True)
    assert _comparable(service.list_applications(db_session)) == _comparable(original)


def test_hebrew_export_uses_original_labels():
    app = Application(submitted_date=dt.date(2026, 1, 2), company="X", title="Y", status="offer",
                      track="ai_rag", source="meta_ai", extra={})
    ws = load_workbook(io.BytesIO(export_workbook([app], "he"))).worksheets[0]
    assert ws.sheet_view.rightToLeft
    assert (ws["I5"].value, ws["F5"].value, ws["E5"].value) == ("הצעה / מתקדם", "Meta Al", "AI Systems & RAG")


# ── CRUD ─────────────────────────────────────────────────────────────────────

def test_create_update_delete_with_validation(ws_session):
    app = service.create(ws_session, {"company": "Acme", "title": "SRE", "status": "submitted",
                                      "submitted_date": dt.date(2026, 10, 1)})
    ws_session.commit()
    uid, first_update = app.row_uid, app.last_update_at
    assert len(uid) == 36 and app.workspace == "personal"

    assert service.update(ws_session, app.id, {"company": "Acme"}) is False       # no change → no timestamp move
    assert service.update(ws_session, app.id, {"status": "interview"}) is True
    ws_session.commit()
    assert app.status == "interview" and app.row_uid == uid and app.last_update_at >= first_update

    for bad in ({"company": ""}, {"status": "hired"}, {"submitted_date": "2026-01-01"}):
        with pytest.raises(service.ValidationError):
            service.update(ws_session, app.id, bad)
    with pytest.raises(service.ValidationError):
        service.create(ws_session, {"company": "OnlyCompany"})

    assert service.delete(ws_session, [app.id]) == 1


def test_setting_a_value_clears_kept_unmapped_import_value(ws_session):
    service.commit_import(ws_session, preview_workbook(_mixed_workbook()))
    gamma = next(a for a in service.list_applications(ws_session) if a.company == "Gamma")
    service.update(ws_session, gamma.id, {"status": "screening"})
    ws_session.commit()
    assert "unmapped" not in (gamma.extra or {})


# ── Statistics (known fixture) ───────────────────────────────────────────────

def _app(status, days_ago, updated_days_ago, track="devops_cloud", source="linkedin", dated=True):
    today = dt.date(2026, 10, 18)
    return Application(
        status=status, track=track, source=source, company="c", title="t", extra={},
        submitted_date=today - dt.timedelta(days=days_ago) if dated else None,
        last_update_at=dt.datetime.combine(today - dt.timedelta(days=updated_days_ago), dt.time(9),
                                           tzinfo=dt.timezone.utc),
    )


def test_kpis_follow_workbook_definitions():
    apps = [
        _app("submitted", 3, 3), _app("screening", 20, 20), _app("initial_call", 10, 13),
        _app("interview", 30, 14), _app("home_assignment", 40, 2), _app("offer", 50, 1),
        _app("rejected", 60, 30), _app("closed_no_response", 70, 40), _app(None, 5, 5),
        _app("submitted", 0, 0, dated=False),
    ]
    s = service.compute_stats(apps, today=dt.date(2026, 10, 18))
    assert s.total_submitted == 9                    # COUNTIF(B,"<>"): rows with a date
    assert s.active == 6                             # submitted×2, screening, initial call, interview, home assignment
    assert s.interviews == 2                         # initial call + interview
    assert s.assignments_offers == 2                 # home assignment + offer
    assert s.by_status["submitted"] == 2 and s.by_status[""] == 1
    # funnel over dated rows: stage reached = index of current status; rejected/closed/unset → submitted only
    assert dict(s.funnel) == {"submitted": 9, "screening": 5, "initial_call": 4, "interview": 3,
                              "home_assignment": 2, "offer": 1}
    # stale: active and last update ≥ 14 days ago (13 days is not stale; 14 is)
    assert sorted(a.status for a in s.stale) == ["interview", "screening"]


def test_weekly_series_fills_gaps():
    apps = [_app("submitted", 0, 0), _app("submitted", 21, 0)]
    weeks = service.compute_stats(apps, today=dt.date(2026, 10, 18)).weekly
    assert [n for _, n in weeks] == [1, 0, 0, 1]
    assert all(w.weekday() == 0 for w, _ in weeks)


def test_filters_apply_to_dates_tracks_sources():
    apps = [_app("submitted", 1, 1, track="ai_rag"), _app("submitted", 40, 1, source="drushim"),
            _app("submitted", 1, 1, dated=False)]
    today = dt.date(2026, 10, 18)
    assert len(service.filter_applications(apps)) == 3
    assert len(service.filter_applications(apps, start=today - dt.timedelta(days=7), end=today)) == 1
    assert len(service.filter_applications(apps, tracks=["ai_rag"])) == 1
    assert len(service.filter_applications(apps, sources=["drushim"])) == 1


def test_label_falls_back_to_code_for_unknown():
    assert label("status", "offer", "he") == "הצעה / מתקדם"
    assert label("status", "weird", "en") == "weird"
    assert label("status", None, "en") == ""
