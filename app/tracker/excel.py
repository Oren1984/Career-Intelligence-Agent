"""Excel import/export with Job_Search_Tracker.xlsx parity.

Import (two steps, nothing is written until ``commit``):
  ``preview_workbook`` reads the tracker sheet, locates the header row by its
  Hebrew or English column names, and classifies every data row:
    - blank rows (template rows with only an ID / formula) are counted and skipped
    - rows without company and job title are reported as "incomplete"
      (imported only when explicitly requested)
    - duplicates (same date, company and title — in the file or already saved) are skipped
    - invalid dates and values outside the dropdown lists are kept verbatim in
      ``extra["unmapped"]``; unknown columns are kept in ``extra["columns"]``
  Issues carry row numbers and field names only, never cell contents.

Export (``export_workbook``) recreates the workbook's three sheets — tracker,
dropdown lists, KPI dashboard — with the same formulas and data validation,
labelled in Hebrew (original layout) or English.
"""
from __future__ import annotations

import datetime as dt
import io
import uuid
from dataclasses import dataclass, field
from typing import Any, Iterable, Optional

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.utils.datetime import from_excel
from openpyxl.worksheet.datavalidation import DataValidation

from app.tracker.vocab import (
    ACTIVE_STATUSES, ASSIGNMENT_OFFER_STATUSES, INTERVIEW_STATUSES, SOURCES, STATUSES, TRACKS,
    VOCABS, label, to_code,
)

# field: (hebrew header — exactly as in the workbook, english header)
COLUMNS: dict[str, tuple[str, str]] = {
    "source_ref": ("מזהה", "ID"),
    "submitted_date": ("תאריך הגשה", "Submission date"),
    "company": ("שם החברה", "Company"),
    "title": ("שם המשרה", "Job title"),
    "track": ("מסלול מקצועי", "Track"),
    "source": ("מקור המשרה", "Source"),
    "location": ("מיקום ומודל עבודה", "Location & work model"),
    "salary_range": ("טווח שכר מוערך", "Estimated salary range"),
    "status": ("סטטוס נוכחי", "Current status"),
    "days": ("ימים מהגשה", "Days since submission"),
    "notes": ("הערות, קישור ואיש קשר", "Notes, link & contact"),
    "closure_reason": ("סיבת סגירה / פידבק", "Closure reason / feedback"),
}
TEXT_FIELDS = ("company", "title", "location", "salary_range", "notes", "closure_reason")
CONTROLLED = ("track", "source", "status")

SHEETS = {
    "tracker": ("מעקב משרות", "Applications"),
    "lists": ("רשימות בחירה", "Lists"),
    "dashboard": ("דשבורד וניתוח", "Dashboard"),
}
_TEXT = {
    "title": ("מעקב הגשת משרות ותהליכי גיוס", "Job applications and hiring processes"),
    "subtitle": ("מסודר כרונולוגית לפי תאריך הגשה", "Ordered chronologically by submission date"),
    "lists_status": ("סטטוסים (Status)", "Statuses"),
    "lists_track": ("מסלולים (Track)", "Tracks"),
    "lists_source": ("מקורות משרה (Source)", "Sources"),
    "dash_title": ("דשבורד מעקב משרות ופעילות גיוס", "Job search dashboard"),
    "dash_subtitle": ("מדדי ביצוע (KPIs), התפלגות סטטוסים ומקורות גיוס - מתעדכן אוטומטית מתוך גיליון 'מעקב משרות'",
                      "KPIs and distributions — calculated automatically from the 'Applications' sheet"),
    "kpi_total": ('סה"כ משרות שהוגשו', "Total submitted"),
    "kpi_active": ("תהליכים פעילים", "Active processes"),
    "kpi_interviews": ("ראיונות שנקבעו", "Interviews scheduled"),
    "kpi_assign": ("מטלות בית והצעות", "Home assignments & offers"),
    "by_status": ("התפלגות לפי סטטוס", "By status"),
    "by_track": ("התפלגות לפי מסלול", "By track"),
    "by_source": ("התפלגות לפי מקור משרה", "By source"),
    "status": ("סטטוס", "Status"), "track": ("מסלול", "Track"), "source": ("מקור", "Source"),
    "count": ("כמות", "Count"),
}
HEADER_ROW = 4
FIRST_DATA_ROW = 5
_DATE_FORMATS = ("%d/%m/%Y", "%d.%m.%Y", "%Y-%m-%d", "%d/%m/%y", "%d-%m-%Y", "%d.%m.%y")


def _t(key: str, lang: str) -> str:
    he, en = _TEXT[key]
    return he if lang == "he" else en


def _norm_header(text) -> str:
    return " ".join(str(text or "").split()).casefold()


_HEADER_LOOKUP = {_norm_header(v): f for f, pair in COLUMNS.items() for v in pair}


# ── Parsing helpers ──────────────────────────────────────────────────────────

def parse_date(value) -> tuple[Optional[dt.date], bool]:
    """Return (date, ok). Empty → (None, True); unparseable → (None, False)."""
    if value is None or (isinstance(value, str) and not value.strip()):
        return None, True
    if isinstance(value, dt.datetime):
        return value.date(), True
    if isinstance(value, dt.date):
        return value, True
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        try:
            converted = from_excel(value)
            return (converted.date() if isinstance(converted, dt.datetime) else converted), True
        except Exception:
            return None, False
    text = str(value).strip()
    for fmt in _DATE_FORMATS:
        try:
            return dt.datetime.strptime(text, fmt).date(), True
        except ValueError:
            continue
    return None, False


def _text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    if isinstance(value, (dt.datetime, dt.date)):
        return value.isoformat()
    return str(value).strip()


def dedup_key(submitted: Optional[dt.date], raw_date: str, company: str, title: str) -> tuple:
    return (submitted.isoformat() if submitted else raw_date.strip(), company.strip().casefold(),
            title.strip().casefold())


# ── Import ───────────────────────────────────────────────────────────────────

@dataclass
class ImportIssue:
    row: int          # worksheet row number
    kind: str         # incomplete | duplicate | invalid_date | unmapped
    field: str = ""


@dataclass
class ImportPreview:
    sheet: str = ""
    rows: list[dict] = field(default_factory=list)            # ready to insert
    incomplete_rows: list[dict] = field(default_factory=list)  # no company and no title
    issues: list[ImportIssue] = field(default_factory=list)
    blank_rows: int = 0
    extra_columns: list[str] = field(default_factory=list)
    error: str = ""

    def count(self, kind: str) -> int:
        return sum(1 for i in self.issues if i.kind == kind)


def _find_header(ws) -> tuple[Optional[int], dict[int, str], dict[int, str]]:
    """Locate the header row: returns (row, {col: field}, {col: unknown header})."""
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 15)):
        known, unknown = {}, {}
        for cell in row:
            if cell.value is None or str(cell.value).strip() == "":
                continue
            fld = _HEADER_LOOKUP.get(_norm_header(cell.value))
            if fld:
                known[cell.column] = fld
            else:
                unknown[cell.column] = str(cell.value).strip()
        if len(known) >= 3 and {"company", "title"} & set(known.values()):
            return row[0].row, known, unknown
    return None, {}, {}


def _pick_sheet(wb):
    names = {_norm_header(n): n for n in wb.sheetnames}
    for candidate in SHEETS["tracker"]:
        if _norm_header(candidate) in names:
            return wb[names[_norm_header(candidate)]]
    return wb.worksheets[0]


def preview_workbook(data: bytes, existing_keys: Iterable[tuple] = ()) -> ImportPreview:
    """Parse an .xlsx file into an ImportPreview without touching the database."""
    preview = ImportPreview()
    try:
        wb = load_workbook(io.BytesIO(data), data_only=True)
    except Exception as exc:
        preview.error = f"Not a readable .xlsx workbook ({type(exc).__name__})."
        return preview
    ws = _pick_sheet(wb)
    preview.sheet = ws.title
    header_row, known, unknown = _find_header(ws)
    if header_row is None:
        preview.error = "No tracker header row found (expected columns such as company and job title)."
        return preview
    preview.extra_columns = list(unknown.values())
    seen = set(existing_keys)

    for row in ws.iter_rows(min_row=header_row + 1, max_row=ws.max_row):
        number = row[0].row
        values = {known[c.column]: c.value for c in row if c.column in known}
        extras = {unknown[c.column]: _text(c.value) for c in row
                  if c.column in unknown and _text(c.value)}
        content_fields = [f for f, v in values.items()
                          if f not in ("source_ref", "days") and _text(v)]
        if not content_fields and not extras:
            preview.blank_rows += 1
            continue

        record: dict[str, Any] = {f: _text(values.get(f)) for f in TEXT_FIELDS}
        record["source_ref"] = _text(values.get("source_ref")) or None
        unmapped: dict[str, str] = {}

        raw_date = values.get("submitted_date")
        submitted, ok = parse_date(raw_date)
        record["submitted_date"] = submitted
        if not ok:
            unmapped["submitted_date"] = _text(raw_date)
            preview.issues.append(ImportIssue(number, "invalid_date", "submitted_date"))

        for fld in CONTROLLED:
            raw = _text(values.get(fld))
            code = to_code(fld, raw)
            record[fld] = code
            if raw and code is None:
                unmapped[fld] = raw
                preview.issues.append(ImportIssue(number, "unmapped", fld))

        record["extra"] = {k: v for k, v in (("unmapped", unmapped), ("columns", extras)) if v}

        if not record["company"] and not record["title"]:
            preview.incomplete_rows.append(record)
            preview.issues.append(ImportIssue(number, "incomplete"))
            continue
        key = dedup_key(submitted, unmapped.get("submitted_date", ""), record["company"], record["title"])
        if key in seen:
            preview.issues.append(ImportIssue(number, "duplicate"))
            continue
        seen.add(key)
        preview.rows.append(record)
    return preview


def new_row_uid() -> str:
    return str(uuid.uuid4())


# ── Export ───────────────────────────────────────────────────────────────────

_HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
_HEADER_FONT = Font(bold=True, color="FFFFFF")


def _controlled_out(app, fld: str, lang: str) -> str:
    code = getattr(app, fld)
    if code:
        return label(fld, code, lang)
    return ((app.extra or {}).get("unmapped") or {}).get(fld, "")


def export_workbook(applications: list, lang: str = "he") -> bytes:
    """Build a tracker workbook (same sheets, dropdowns and formulas as the original)."""
    lang = "he" if lang == "he" else "en"
    wb = Workbook()
    ws = wb.active
    ws.title = SHEETS["tracker"][0 if lang == "he" else 1]
    lists = wb.create_sheet(SHEETS["lists"][0 if lang == "he" else 1])
    dash = wb.create_sheet(SHEETS["dashboard"][0 if lang == "he" else 1])
    for sheet in (ws, lists, dash):
        sheet.sheet_view.rightToLeft = lang == "he"

    extra_headers: list[str] = []
    for app in applications:
        for name in ((app.extra or {}).get("columns") or {}):
            if name not in extra_headers:
                extra_headers.append(name)

    ws["A1"] = _t("title", lang)
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = _t("subtitle", lang)
    headers = [pair[0 if lang == "he" else 1] for pair in COLUMNS.values()] + extra_headers
    for col, text in enumerate(headers, start=1):
        cell = ws.cell(row=HEADER_ROW, column=col, value=text)
        cell.fill, cell.font = _HEADER_FILL, _HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for i, app in enumerate(applications):
        r = FIRST_DATA_ROW + i
        raw_date = ((app.extra or {}).get("unmapped") or {}).get("submitted_date")
        row = [
            app.source_ref or i + 1,
            app.submitted_date or raw_date or None,
            app.company, app.title,
            _controlled_out(app, "track", lang), _controlled_out(app, "source", lang),
            app.location, app.salary_range,
            _controlled_out(app, "status", lang),
            f'=IF(B{r}="","",IFERROR(DATEDIF(B{r},TODAY(),"d"),IFERROR(TODAY()-B{r},"")))',
            app.notes, app.closure_reason,
        ] + [((app.extra or {}).get("columns") or {}).get(h, "") for h in extra_headers]
        for col, value in enumerate(row, start=1):
            ws.cell(row=r, column=col, value=value if value != "" else None)
        ws.cell(row=r, column=2).number_format = "DD/MM/YYYY"

    last = max(FIRST_DATA_ROW + len(applications) + 50, 59)
    for col, width in zip("ABCDEFGHIJKL", (7, 13, 22, 26, 22, 22, 22, 16, 16, 10, 40, 32)):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = f"A{FIRST_DATA_ROW}"

    # Dropdown lists (sheet 2) and data validation, as in the original workbook.
    for col, (key, vocab) in enumerate((("lists_status", STATUSES), ("lists_track", TRACKS),
                                        ("lists_source", SOURCES)), start=1):
        lists.cell(row=1, column=col, value=_t(key, lang)).font = Font(bold=True)
        for j, code in enumerate(vocab, start=2):
            lists.cell(row=j, column=col, value=label({1: "status", 2: "track", 3: "source"}[col], code, lang))
        lists.column_dimensions[get_column_letter(col)].width = 30
    sheet_ref = f"'{lists.title}'"
    for col_letter, list_col, size in (("E", "B", len(TRACKS)), ("F", "C", len(SOURCES)), ("I", "A", len(STATUSES))):
        dv = DataValidation(type="list", formula1=f"={sheet_ref}!${list_col}$2:${list_col}${size + 1}", allow_blank=True)
        dv.add(f"{col_letter}{FIRST_DATA_ROW}:{col_letter}{last}")
        ws.add_data_validation(dv)

    _write_dashboard(dash, ws.title, lang, last)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _write_dashboard(dash, tracker_title: str, lang: str, last: int) -> None:
    src = f"'{tracker_title}'"
    rng = lambda col: f"{src}!${col}${FIRST_DATA_ROW}:${col}${last}"  # noqa: E731
    dash["A1"] = _t("dash_title", lang)
    dash["A1"].font = Font(bold=True, size=14)
    dash["A2"] = _t("dash_subtitle", lang)
    status_label = lambda code: label("status", code, lang)  # noqa: E731
    countifs = lambda codes: " + ".join(f'COUNTIF({rng("I")}, "{status_label(c)}")' for c in codes)  # noqa: E731
    kpis = (("A", "kpi_total", f'=COUNTIF({rng("B")}, "<>")'),
            ("C", "kpi_active", "=" + countifs(ACTIVE_STATUSES)),
            ("E", "kpi_interviews", "=" + countifs(INTERVIEW_STATUSES)),
            ("G", "kpi_assign", "=" + countifs(ASSIGNMENT_OFFER_STATUSES)))
    for col, key, formula in kpis:
        dash[f"{col}4"] = _t(key, lang)
        dash[f"{col}4"].font = Font(bold=True)
        dash[f"{col}5"] = formula
    for start_col, title_key, head_key, fld, data_col in (
        ("A", "by_status", "status", "status", "I"), ("D", "by_track", "track", "track", "E"),
        ("G", "by_source", "source", "source", "F"),
    ):
        c0 = ord(start_col) - 64
        dash.cell(row=8, column=c0, value=_t(title_key, lang)).font = Font(bold=True)
        dash.cell(row=9, column=c0, value=_t(head_key, lang)).font = Font(bold=True)
        dash.cell(row=9, column=c0 + 1, value=_t("count", lang)).font = Font(bold=True)
        for i, code in enumerate(VOCABS[fld], start=10):
            name_cell = dash.cell(row=i, column=c0, value=label(fld, code, lang))
            dash.cell(row=i, column=c0 + 1, value=f"=COUNTIF({rng(data_col)}, {name_cell.coordinate})")
    for col in "ABCDEFGH":
        dash.column_dimensions[col].width = 24
